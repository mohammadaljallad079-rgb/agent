"""Format parsers for the document ingestion layer.

Each parser turns a file into a :class:`~arabic_agent.documents.models.Document`
of ordered, provenance-bearing sections. Parsers are:

* **bounded** — page/sheet/slide/row/section/size caps so a huge file cannot
  exhaust memory or the model context;
* **honest** — when a parser dependency is missing the format is reported as
  *unsupported*, and when a document yields no text (a scanned PDF) the
  ``has_text`` flag is False rather than pretending extraction succeeded;
* **lazy** — parser libraries are imported inside the function so the platform
  runs without them; only the formats whose dependency is installed work.

The registry is the single source of truth for "which formats can this build
actually read", which the capability catalogue and the tools both consult.
"""

from __future__ import annotations

import csv
import io
import json
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from ..context.budget import estimate_tokens
from ..core.errors import ValidationError
from .models import Document, DocumentMetadata, DocumentSection, Provenance

#: Formats the layer understands, mapped to a canonical name.
_EXTENSION_FORMATS: dict[str, str] = {
    ".txt": "text",
    ".log": "text",
    ".rst": "text",
    ".ini": "text",
    ".cfg": "text",
    ".conf": "text",
    ".toml": "text",
    ".yaml": "text",
    ".yml": "text",
    ".md": "markdown",
    ".markdown": "markdown",
    ".json": "json",
    ".jsonl": "json",
    ".csv": "csv",
    ".tsv": "csv",
    ".html": "html",
    ".htm": "html",
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".pptx": "pptx",
}


@dataclass(frozen=True)
class ParseLimits:
    """Caps applied to every parse so nothing is unbounded."""

    max_sections: int = 2000
    max_section_chars: int = 20_000
    max_pages: int = 500
    max_sheets: int = 50
    max_rows_per_sheet: int = 5000
    max_cols_per_sheet: int = 100
    max_slides: int = 500
    max_text_bytes: int = 8 * 1024 * 1024


def format_for(path: Path) -> str:
    """Canonical format name for a path, or ``"unknown"``."""
    return _EXTENSION_FORMATS.get(path.suffix.lower(), "unknown")


def _section(
    *,
    path: str,
    index: int,
    kind: str,
    text: str,
    locator: str,
    limits: ParseLimits,
    level: int | None = None,
    **provenance_kwargs: Any,
) -> DocumentSection:
    capped = text[: limits.max_section_chars]
    return DocumentSection(
        id=f"{path}#{index}",
        kind=kind,
        text=capped,
        level=level,
        provenance=Provenance(path=path, locator=locator, **provenance_kwargs),
        tokens=estimate_tokens(capped),
    )


def parse_text(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    raw = path.read_bytes()
    truncated = len(raw) > limits.max_text_bytes
    text = raw[: limits.max_text_bytes].decode("utf-8", errors="replace")
    lines = text.splitlines()
    sections: list[DocumentSection] = []
    chunk: list[str] = []
    start_line = 1
    for number, line in enumerate(lines, start=1):
        chunk.append(line)
        if len(chunk) >= 200 or number == len(lines):
            block = "\n".join(chunk)
            sections.append(
                _section(
                    path=display_path,
                    index=len(sections),
                    kind="text",
                    text=block,
                    locator=f"lines {start_line}-{number}",
                    limits=limits,
                    line=start_line,
                    end_line=number,
                )
            )
            chunk = []
            start_line = number + 1
        if len(sections) >= limits.max_sections:
            truncated = True
            break
    return Document(
        path=display_path,
        format="text",
        metadata=DocumentMetadata(row_count=len(lines)),
        sections=tuple(sections),
        truncated=truncated,
        has_text=bool(text.strip()),
    )


def parse_markdown(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    raw = path.read_bytes()
    truncated = len(raw) > limits.max_text_bytes
    text = raw[: limits.max_text_bytes].decode("utf-8", errors="replace")
    lines = text.splitlines()
    sections: list[DocumentSection] = []
    buffer: list[str] = []
    start_line = 1
    level = 0
    heading_title: str | None = None

    def flush(end_line: int) -> None:
        if not buffer:
            return
        block = "\n".join(buffer).strip("\n")
        if block:
            kind = "heading" if level else "text"
            locator = f"lines {start_line}-{end_line}"
            sections.append(
                _section(
                    path=display_path,
                    index=len(sections),
                    kind=kind,
                    text=block,
                    locator=locator,
                    limits=limits,
                    level=level or None,
                    line=start_line,
                    end_line=end_line,
                )
            )

    for number, line in enumerate(lines, start=1):
        stripped = line.lstrip()
        if stripped.startswith("#"):
            flush(number - 1)
            hashes = len(stripped) - len(stripped.lstrip("#"))
            level = min(hashes, 6)
            buffer = [line]
            start_line = number
            if heading_title is None:
                heading_title = stripped.lstrip("#").strip() or None
        else:
            if not buffer:
                start_line = number
            buffer.append(line)
        if len(sections) >= limits.max_sections:
            truncated = True
            break
    flush(len(lines))
    return Document(
        path=display_path,
        format="markdown",
        metadata=DocumentMetadata(title=heading_title),
        sections=tuple(sections[: limits.max_sections]),
        truncated=truncated,
        has_text=bool(text.strip()),
    )


def parse_json(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    raw = path.read_bytes()
    truncated = len(raw) > limits.max_text_bytes
    text = raw[: limits.max_text_bytes].decode("utf-8", errors="replace")
    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValidationError("invalid JSON document", details={"error": str(exc)}) from exc
    pretty = json.dumps(payload, ensure_ascii=False, indent=2)
    sections: list[DocumentSection] = []
    step = limits.max_section_chars
    for offset in range(0, len(pretty), step):
        sections.append(
            _section(
                path=display_path,
                index=len(sections),
                kind="text",
                text=pretty[offset : offset + step],
                locator=f"chars {offset}-{offset + step}",
                limits=limits,
            )
        )
        if len(sections) >= limits.max_sections:
            truncated = True
            break
    return Document(
        path=display_path,
        format="json",
        sections=tuple(sections),
        truncated=truncated,
        has_text=bool(pretty.strip()),
    )


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return "\t" if "\t" in sample.splitlines()[0] else ","


def parse_csv(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    raw = path.read_bytes()
    truncated = len(raw) > limits.max_text_bytes
    text = raw[: limits.max_text_bytes].decode("utf-8", errors="replace")
    sample = "\n".join(text.splitlines()[:5])
    delimiter = _sniff_delimiter(sample) if sample.strip() else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    header = rows[0] if rows else []
    sections: list[DocumentSection] = []
    if header:
        sections.append(
            _section(
                path=display_path,
                index=0,
                kind="table",
                text=f"columns: {', '.join(header)}",
                locator="header row",
                limits=limits,
                range="row 1",
            )
        )
    body = rows[1:]
    if len(body) > limits.max_rows_per_sheet:
        truncated = True
        body = body[: limits.max_rows_per_sheet]
    block: list[str] = []
    start_row = 2
    for number, row in enumerate(body, start=2):
        block.append(delimiter.join(row))
        if len(block) >= 100 or number == len(body) + 1:
            sections.append(
                _section(
                    path=display_path,
                    index=len(sections),
                    kind="row",
                    text="\n".join(block),
                    locator=f"rows {start_row}-{number}",
                    limits=limits,
                    range=f"rows {start_row}-{number}",
                )
            )
            block = []
            start_row = number + 1
        if len(sections) >= limits.max_sections:
            truncated = True
            break
    return Document(
        path=display_path,
        format="csv",
        metadata=DocumentMetadata(
            row_count=len(rows),
            column_count=len(header),
            extra={"delimiter": delimiter},
        ),
        sections=tuple(sections),
        truncated=truncated,
        has_text=bool(text.strip()),
    )


def parse_html(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    try:
        from bs4 import BeautifulSoup
    except ImportError as exc:  # pragma: no cover - dependency presence varies
        raise ValidationError("HTML parsing requires beautifulsoup4 (extras: documents)") from exc
    raw = path.read_bytes()
    truncated = len(raw) > limits.max_text_bytes
    soup = BeautifulSoup(
        raw[: limits.max_text_bytes].decode("utf-8", errors="replace"), "html.parser"
    )
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    sections: list[DocumentSection] = []
    title = soup.title.string.strip() if soup.title and soup.title.string else None
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th"]):
        content = element.get_text(" ", strip=True)
        if not content:
            continue
        name = element.name
        kind = "heading" if name.startswith("h") else "text"
        level = int(name[1]) if kind == "heading" else None
        sections.append(
            _section(
                path=display_path,
                index=len(sections),
                kind=kind,
                text=content,
                locator=f"{name} #{len(sections) + 1}",
                limits=limits,
                level=level,
            )
        )
        if len(sections) >= limits.max_sections:
            truncated = True
            break
    if not sections:
        body_text = soup.get_text("\n", strip=True)
        if body_text:
            sections.append(
                _section(
                    path=display_path,
                    index=0,
                    kind="text",
                    text=body_text,
                    locator="body",
                    limits=limits,
                )
            )
    return Document(
        path=display_path,
        format="html",
        metadata=DocumentMetadata(title=title),
        sections=tuple(sections),
        truncated=truncated,
        has_text=any(s.text.strip() for s in sections),
    )


def parse_pdf(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("PDF parsing requires pypdf (extras: documents)") from exc
    reader = PdfReader(str(path))
    page_count = len(reader.pages)
    metadata = DocumentMetadata(page_count=page_count)
    warnings: list[str] = []
    try:
        info = reader.metadata
        if info is not None:
            metadata = DocumentMetadata(
                page_count=page_count,
                title=str(info.title) if info.title else None,
                author=str(info.author) if info.author else None,
                created=str(info.creation_date) if info.creation_date else None,
            )
    except Exception:
        warnings.append("pdf metadata could not be read")

    sections: list[DocumentSection] = []
    truncated = False
    pages_to_read = min(page_count, limits.max_pages)
    if page_count > limits.max_pages:
        truncated = True
        warnings.append(f"only the first {limits.max_pages} of {page_count} pages were read")
    for page_index in range(pages_to_read):
        try:
            text = reader.pages[page_index].extract_text() or ""
        except Exception as exc:
            warnings.append(f"page {page_index + 1} could not be read: {type(exc).__name__}")
            continue
        if not text.strip():
            continue
        sections.append(
            _section(
                path=display_path,
                index=len(sections),
                kind="page",
                text=text,
                locator=f"page {page_index + 1}",
                limits=limits,
                page=page_index + 1,
            )
        )
        if len(sections) >= limits.max_sections:
            truncated = True
            break
    has_text = any(s.text.strip() for s in sections)
    if not has_text:
        warnings.append(
            "no extractable text found; the PDF is likely scanned images. "
            "OCR is not implemented, so text extraction did not succeed."
        )
    return Document(
        path=display_path,
        format="pdf",
        metadata=metadata,
        sections=tuple(sections),
        truncated=truncated,
        has_text=has_text,
        warnings=tuple(warnings),
    )


def parse_docx(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("DOCX parsing requires python-docx (extras: documents)") from exc
    document = docx.Document(str(path))
    sections: list[DocumentSection] = []
    truncated = False
    paragraph_count = 0
    for paragraph in document.paragraphs:
        content = paragraph.text.strip()
        if not content:
            continue
        paragraph_count += 1
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        kind = "text"
        level: int | None = None
        if style.startswith("heading"):
            kind = "heading"
            digits = "".join(ch for ch in style if ch.isdigit())
            level = int(digits) if digits else 1
        elif style == "title":
            kind = "heading"
            level = 1
        sections.append(
            _section(
                path=display_path,
                index=len(sections),
                kind=kind,
                text=content,
                locator=f"paragraph {paragraph_count}",
                limits=limits,
                level=level,
            )
        )
        if len(sections) >= limits.max_sections:
            truncated = True
            break
    table_count = 0
    if not truncated:
        for table in document.tables:
            table_count += 1
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(" | ".join(cells))
            sections.append(
                _section(
                    path=display_path,
                    index=len(sections),
                    kind="table",
                    text="\n".join(rows),
                    locator=f"table {table_count}",
                    limits=limits,
                )
            )
            if len(sections) >= limits.max_sections:
                truncated = True
                break
    props = getattr(document, "core_properties", None)
    return Document(
        path=display_path,
        format="docx",
        metadata=DocumentMetadata(
            paragraph_count=paragraph_count,
            title=(props.title if props and props.title else None),
            author=(props.author if props and props.author else None),
            extra={"table_count": table_count},
        ),
        sections=tuple(sections),
        truncated=truncated,
        has_text=bool(sections),
    )


def parse_xlsx(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("XLSX parsing requires openpyxl (extras: documents)") from exc
    workbook = load_workbook(str(path), read_only=True, data_only=False)
    try:
        sheet_names = list(workbook.sheetnames)
        sections: list[DocumentSection] = []
        warnings: list[str] = []
        truncated = False
        for sheet_index, sheet_name in enumerate(sheet_names):
            if sheet_index >= limits.max_sheets:
                truncated = True
                warnings.append(f"only the first {limits.max_sheets} sheets were read")
                break
            worksheet = workbook[sheet_name]
            rows_written = 0
            block: list[str] = []
            start_row = 1
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if row_index > limits.max_rows_per_sheet:
                    truncated = True
                    warnings.append(
                        f"sheet {sheet_name!r} truncated at {limits.max_rows_per_sheet} rows"
                    )
                    break
                cells = [
                    "" if value is None else str(value)
                    for value in row[: limits.max_cols_per_sheet]
                ]
                if not any(cells):
                    continue
                block.append(" | ".join(cells))
                rows_written += 1
                if len(block) >= 100:
                    sections.append(
                        _section(
                            path=display_path,
                            index=len(sections),
                            kind="sheet",
                            text="\n".join(block),
                            locator=f"sheet {sheet_name!r} rows {start_row}-{row_index}",
                            limits=limits,
                            sheet=sheet_name,
                            range=f"rows {start_row}-{row_index}",
                        )
                    )
                    block = []
                    start_row = row_index + 1
                if len(sections) >= limits.max_sections:
                    truncated = True
                    break
            if block:
                sections.append(
                    _section(
                        path=display_path,
                        index=len(sections),
                        kind="sheet",
                        text="\n".join(block),
                        locator=f"sheet {sheet_name!r} rows {start_row}-end",
                        limits=limits,
                        sheet=sheet_name,
                        range=f"rows {start_row}-end",
                    )
                )
            if len(sections) >= limits.max_sections:
                truncated = True
                break
        return Document(
            path=display_path,
            format="xlsx",
            metadata=DocumentMetadata(
                sheet_count=len(sheet_names),
                extra={"sheets": sheet_names[: limits.max_sheets]},
            ),
            sections=tuple(sections),
            truncated=truncated,
            has_text=bool(sections),
            warnings=tuple(warnings),
        )
    finally:
        workbook.close()


def parse_pptx(path: Path, *, display_path: str, limits: ParseLimits) -> Document:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("PPTX parsing requires python-pptx (extras: documents)") from exc
    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    slide_count = len(slides)
    sections: list[DocumentSection] = []
    truncated = False
    for slide_index, slide in enumerate(slides):
        if slide_index >= limits.max_slides:
            truncated = True
            break
        texts: list[str] = []
        title: str | None = None
        for shape in slide.shapes:
            if shape.has_text_frame:
                content = shape.text_frame.text.strip()
                if content:
                    texts.append(content)
                    if title is None and getattr(shape, "is_placeholder", False):
                        try:
                            if shape.placeholder_format.idx == 0:
                                title = content
                        except (AttributeError, ValueError):
                            pass
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = []
                for row in shape.table.rows:
                    rows.append(" | ".join(cell.text.strip() for cell in row.cells))
                if rows:
                    texts.append("\n".join(rows))
        if texts:
            sections.append(
                _section(
                    path=display_path,
                    index=len(sections),
                    kind="slide",
                    text="\n".join(texts),
                    locator=f"slide {slide_index + 1}",
                    limits=limits,
                    slide=slide_index + 1,
                )
            )
        notes = _pptx_notes(slide)
        if notes:
            sections.append(
                _section(
                    path=display_path,
                    index=len(sections),
                    kind="note",
                    text=notes,
                    locator=f"slide {slide_index + 1} notes",
                    limits=limits,
                    slide=slide_index + 1,
                )
            )
        if len(sections) >= limits.max_sections:
            truncated = True
            break
    return Document(
        path=display_path,
        format="pptx",
        metadata=DocumentMetadata(slide_count=slide_count),
        sections=tuple(sections),
        truncated=truncated,
        has_text=bool(sections),
    )


def _pptx_notes(slide: object) -> str | None:
    """Speaker notes, if safely accessible. Never fatal when absent."""
    try:
        notes_slide = slide.notes_slide  # type: ignore[attr-defined]
        text = notes_slide.notes_text_frame.text.strip()
        return text or None
    except Exception:
        return None


#: Parser dispatch. Format name -> callable.
PARSERS: dict[str, Callable[..., Document]] = {
    "text": parse_text,
    "markdown": parse_markdown,
    "json": parse_json,
    "csv": parse_csv,
    "html": parse_html,
    "pdf": parse_pdf,
    "docx": parse_docx,
    "xlsx": parse_xlsx,
    "pptx": parse_pptx,
}

#: Format -> the importable module that must be present for it to work. Formats
#: absent from this map have no third-party dependency (stdlib only).
_REQUIRED_MODULES: dict[str, str] = {
    "html": "bs4",
    "pdf": "pypdf",
    "docx": "docx",
    "xlsx": "openpyxl",
    "pptx": "pptx",
}


def supported_formats() -> tuple[str, ...]:
    """Formats this build can actually parse right now (dependency present)."""
    import importlib.util

    supported: list[str] = []
    for name in PARSERS:
        module = _REQUIRED_MODULES.get(name)
        if module is None or importlib.util.find_spec(module) is not None:
            supported.append(name)
    return tuple(sorted(supported))


def is_supported(fmt: str) -> bool:
    return fmt in supported_formats()
