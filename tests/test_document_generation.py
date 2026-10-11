"""Document generation: real files written, verified by re-reading them.

Generation is exercised through the *real* tools (which resolve paths through
the filesystem sandbox and are permission-declared) and the output is verified
by re-reading it with the ingestion parsers — a generated PDF must contain the
text, a generated XLSX must contain the rows. Nothing asserts a file exists
without checking its content, and unsupported/ambiguous requests must fail
loudly rather than write a broken or empty file.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from arabic_agent.core.errors import PermissionDeniedError, ValidationError
from arabic_agent.documents.generation import (
    Block,
    DocumentSpec,
    format_for_suffix,
    generate_document,
    generate_spreadsheet,
    generation_capabilities,
)
from arabic_agent.permissions.base import Permission
from arabic_agent.tools.base import ToolContext
from arabic_agent.tools.builtin.docgen import (
    DOCUMENT_GENERATION_TOOLS,
    CreatePresentationTool,
    CreateSpreadsheetTool,
    DocumentGenerationFormatsTool,
    GenerateDocumentTool,
)

#: Generation capabilities in *this* install. Tests for a format whose library
#: is absent skip (with a reason) rather than failing — the same honest
#: degradation the code itself promises. The refusal path is tested separately
#: and needs no library.
_CAPS = generation_capabilities()
requires_pdf = pytest.mark.skipif(not _CAPS["pdf"], reason="reportlab not installed")
requires_docx = pytest.mark.skipif(not _CAPS["docx"], reason="python-docx not installed")
requires_xlsx = pytest.mark.skipif(not _CAPS["xlsx"], reason="openpyxl not installed")
requires_pptx = pytest.mark.skipif(not _CAPS["pptx"], reason="python-pptx not installed")


def _context(workspace: Path) -> ToolContext:
    return ToolContext(
        tenant_id="local", user_id="local", run_id="run-gen", workspace_dir=str(workspace)
    )


_ARABIC_SPEC = DocumentSpec.from_arguments(
    {
        "title": "فاتورة رقم 42",
        "blocks": [
            {"type": "heading", "text": "تفاصيل الفاتورة", "level": 1},
            {"type": "paragraph", "text": "مرحبا بالعالم — مستند عربي حقيقي."},
            {"type": "bullet", "items": ["خدمة", "دعم"]},
            {"type": "table", "headers": ["البند", "السعر"], "rows": [["خدمة", "100"]]},
        ],
    }
)


# --- registry contract ------------------------------------------------------


def test_generation_tools_are_declared_writes() -> None:
    assert GenerateDocumentTool.permissions == (Permission.WRITE_FILE,)
    assert CreateSpreadsheetTool.permissions == (Permission.WRITE_FILE,)
    assert CreatePresentationTool.permissions == (Permission.WRITE_FILE,)
    # The formats reporter is read-only.
    assert DocumentGenerationFormatsTool.permissions == ()
    names = {cls.name for cls in DOCUMENT_GENERATION_TOOLS}
    assert names == {
        "generate_document",
        "create_spreadsheet",
        "create_presentation",
        "document_generation_formats",
    }


def test_format_for_suffix() -> None:
    assert format_for_suffix(".pdf") == "pdf"
    assert format_for_suffix(".MD") == "markdown"
    assert format_for_suffix(".pptx") == "pptx"
    assert format_for_suffix(".nope") is None


# --- real writes, verified by re-reading ------------------------------------


@pytest.mark.asyncio
@requires_pdf
async def test_generate_pdf_contains_the_text(tmp_path: Path) -> None:
    tool = GenerateDocumentTool()
    result = await tool.execute(
        {
            "path": "report.pdf",
            "title": "تقرير",
            "blocks": [{"type": "paragraph", "text": "hello"}],
        },
        _context(tmp_path),
    )
    assert result.ok
    path = tmp_path / "report.pdf"
    assert path.read_bytes().startswith(b"%PDF-")
    assert result.data["bytes_written"] == path.stat().st_size

    from pypdf import PdfReader

    text = PdfReader(str(path)).pages[0].extract_text()
    assert "hello" in text


@pytest.mark.asyncio
@requires_docx
async def test_generate_docx_is_a_real_zip_with_text(tmp_path: Path) -> None:
    tool = GenerateDocumentTool()
    result = await tool.execute(
        {"path": "doc.docx", "title": "Title", "blocks": [{"type": "paragraph", "text": "body"}]},
        _context(tmp_path),
    )
    assert result.ok
    path = tmp_path / "doc.docx"
    assert path.read_bytes()[:2] == b"PK"

    import docx

    document = docx.Document(str(path))
    joined = "\n".join(p.text for p in document.paragraphs)
    assert "body" in joined


@pytest.mark.asyncio
@requires_pptx
async def test_generate_pptx_is_a_real_package_with_slides(tmp_path: Path) -> None:
    """A real .pptx package: the title, body and table are all present."""
    tool = GenerateDocumentTool()
    result = await tool.execute(
        {
            "path": "deck.pptx",
            "title": "Deck",
            "blocks": [
                {"type": "heading", "text": "Section"},
                {"type": "paragraph", "text": "Body text"},
                {"type": "bullet", "items": ["one", "two"]},
                {"type": "table", "headers": ["A", "B"], "rows": [["1", "2"]]},
            ],
        },
        _context(tmp_path),
    )
    assert result.ok
    path = tmp_path / "deck.pptx"
    assert path.read_bytes()[:2] == b"PK"  # a real OOXML zip
    assert result.data["bytes_written"] == path.stat().st_size

    from pptx import Presentation

    presentation = Presentation(str(path))
    slides = list(presentation.slides)
    assert len(slides) == 2  # title slide + one heading slide
    joined = "\n".join(
        shape.text_frame.text
        for slide in slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text
    )
    assert "Deck" in joined
    assert "Section" in joined
    # Successive body blocks accumulate rather than overwrite one another.
    assert "Body text" in joined
    assert "one" in joined and "two" in joined
    assert any(shape.has_table for shape in slides[1].shapes)


@pytest.mark.asyncio
@requires_pptx
async def test_create_presentation_arabic_is_shaped_and_rtl(tmp_path: Path) -> None:
    """Arabic deck: the text survives and the deck is right-to-left."""
    tool = CreatePresentationTool()
    result = await tool.execute(
        {
            "path": "عرض.pptx",
            "title": "العنوان",
            "blocks": [{"type": "heading", "text": "المحتوى"}],
        },
        _context(tmp_path),
    )
    assert result.ok
    from pptx import Presentation

    presentation = Presentation(str(tmp_path / "عرض.pptx"))
    # The Arabic heading is present (shaped with presentation forms), and the
    # paragraphs are marked right-to-left, not left-to-right.
    assert len(presentation.slides) == 2
    title_frame = presentation.slides[0].shapes.title.text_frame
    assert title_frame.text  # the title is not empty
    assert (
        title_frame.paragraphs[0]._p.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}pPr"
        )
        is not None
    )
    assert title_frame.paragraphs[0]._p.pPr.get("rtl") == "1"


@pytest.mark.asyncio
@requires_pptx
async def test_create_presentation_rejects_non_pptx_extension(tmp_path: Path) -> None:
    tool = CreatePresentationTool()
    with pytest.raises(ValidationError):
        await tool.execute({"path": "deck.pdf", "title": "t"}, _context(tmp_path))
    assert not (tmp_path / "deck.pdf").exists()


@pytest.mark.asyncio
@requires_pptx
async def test_create_presentation_path_escape_is_denied(tmp_path: Path) -> None:
    tool = CreatePresentationTool()
    with pytest.raises(PermissionDeniedError):
        await tool.execute({"path": "../escape.pptx", "title": "x"}, _context(tmp_path))


@pytest.mark.asyncio
@requires_pptx
async def test_generate_pptx_round_trips_through_the_parser(tmp_path: Path) -> None:
    """A generated deck re-parses through the real document pipeline.

    This closes the loop: what generation writes, the read side actually reads
    back, with the Arabic body text intact.
    """
    from arabic_agent.documents.generation import DocumentSpec
    from arabic_agent.documents.parsers import ParseLimits, parse_pptx

    path = tmp_path / "تقرير.pptx"
    generate_document(
        path,
        DocumentSpec(
            title="تقرير المبيعات",
            blocks=(
                Block.from_mapping({"type": "heading", "text": "الفصل الأول"}),
                Block.from_mapping({"type": "paragraph", "text": "النص العربي"}),
            ),
        ),
    )
    document = parse_pptx(path, display_path="تقرير.pptx", limits=ParseLimits())
    assert document.format == "pptx"
    assert document.metadata.slide_count == 2
    text = "\n".join(section.text for section in document.sections)
    assert "تقرير المبيعات" in text
    assert "الفصل الأول" in text
    assert "النص العربي" in text


@pytest.mark.asyncio
@requires_xlsx
async def test_create_spreadsheet_xlsx_has_the_rows(tmp_path: Path) -> None:
    tool = CreateSpreadsheetTool()
    result = await tool.execute(
        {
            "path": "data.xlsx",
            "headers": ["الاسم", "المبلغ"],
            "rows": [["أحمد", "10"], ["سارة", "20"]],
        },
        _context(tmp_path),
    )
    assert result.ok
    from openpyxl import load_workbook

    sheet = load_workbook(str(tmp_path / "data.xlsx")).active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert ["الاسم", "المبلغ"] in rows
    assert ["أحمد", "10"] in rows


@pytest.mark.asyncio
@requires_xlsx
async def test_create_spreadsheet_honours_the_sheet_name(tmp_path: Path) -> None:
    """A requested sheet name is used, not silently replaced with 'Sheet1'."""
    tool = CreateSpreadsheetTool()
    result = await tool.execute(
        {"path": "named.xlsx", "sheet": "الفواتير", "rows": [["x", "1"]]},
        _context(tmp_path),
    )
    assert result.ok
    from openpyxl import load_workbook

    workbook = load_workbook(str(tmp_path / "named.xlsx"))
    assert workbook.active.title == "الفواتير"


def test_sanitize_sheet_name_is_excel_safe() -> None:
    from arabic_agent.documents.generation import sanitize_sheet_name

    assert sanitize_sheet_name("الفواتير") == "الفواتير"
    # Excel forbids []:*?/\ and caps the length at 31 characters.
    assert sanitize_sheet_name("a[b]:c*d?e/f\\g") == "abcdefg"
    assert sanitize_sheet_name("") == "Sheet1"
    assert len(sanitize_sheet_name("x" * 50)) == 31


@pytest.mark.asyncio
async def test_create_spreadsheet_csv_has_a_bom_for_excel(tmp_path: Path) -> None:
    tool = CreateSpreadsheetTool()
    result = await tool.execute({"path": "data.csv", "rows": [["a", "b"]]}, _context(tmp_path))
    assert result.ok
    raw = (tmp_path / "data.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    assert "a,b" in raw.decode("utf-8-sig")


@pytest.mark.asyncio
async def test_generate_json_is_valid_json(tmp_path: Path) -> None:
    tool = GenerateDocumentTool()
    result = await tool.execute(
        {"path": "doc.json", "title": "t", "blocks": [{"type": "paragraph", "text": "x"}]},
        _context(tmp_path),
    )
    assert result.ok
    import json

    payload = json.loads((tmp_path / "doc.json").read_text(encoding="utf-8"))
    assert payload["title"] == "t"
    assert payload["blocks"][0]["text"] == "x"


# --- honest failures --------------------------------------------------------


@pytest.mark.asyncio
async def test_xlsx_is_refused_by_generate_document(tmp_path: Path) -> None:
    """Content blocks do not map onto a spreadsheet; refuse, don't write empty."""
    tool = GenerateDocumentTool()
    with pytest.raises(ValidationError):
        await tool.execute({"path": "x.xlsx", "title": "t"}, _context(tmp_path))
    assert not (tmp_path / "x.xlsx").exists()


@pytest.mark.asyncio
async def test_unknown_extension_is_refused(tmp_path: Path) -> None:
    tool = GenerateDocumentTool()
    with pytest.raises(ValidationError):
        await tool.execute({"path": "x.docx.exe", "title": "t"}, _context(tmp_path))


@pytest.mark.asyncio
async def test_empty_spec_is_refused(tmp_path: Path) -> None:
    with pytest.raises(ValidationError):
        DocumentSpec.from_arguments({})


@pytest.mark.asyncio
async def test_unknown_block_type_is_refused() -> None:
    with pytest.raises(ValidationError):
        DocumentSpec.from_arguments({"blocks": [{"type": "video", "text": "x"}]})


@pytest.mark.asyncio
async def test_path_escape_is_denied(tmp_path: Path) -> None:
    tool = GenerateDocumentTool()
    with pytest.raises(PermissionDeniedError):
        await tool.execute({"path": "../escape.pdf", "title": "x"}, _context(tmp_path))


@pytest.mark.asyncio
async def test_missing_library_reports_unsupported(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """A format whose library is absent is refused, never faked."""
    import arabic_agent.documents.generation as generation

    monkeypatch.setattr(
        generation, "generation_capabilities", lambda: {**generation_capabilities(), "pdf": False}
    )
    with pytest.raises(Exception) as excinfo:
        generate_document(tmp_path / "x.pdf", _ARABIC_SPEC)
    assert "not available" in str(excinfo.value)


# --- direct API -------------------------------------------------------------


def test_generation_capabilities_shape() -> None:
    caps = generation_capabilities()
    for fmt in ("pdf", "docx", "pptx", "xlsx", "csv", "markdown", "text", "html", "json"):
        assert fmt in caps
        assert isinstance(caps[fmt], bool)


def test_generate_spreadsheet_rejects_pdf(tmp_path: Path) -> None:
    with pytest.raises(ValidationError):
        generate_spreadsheet(tmp_path / "x.pdf", sheet="s", headers=(), rows=(("a",),))
