"""Phase 6: document ingestion, provenance, retrieval, archives and secrets.

Documents are generated for real with the parser libraries (python-docx,
openpyxl, python-pptx) and a small hand-built PDF, then parsed by the *real*
ingestion layer. Nothing here fabricates extraction: if a format cannot be
parsed the test asserts the honest failure instead.
"""

from __future__ import annotations

import io
import tarfile
import zipfile
from pathlib import Path

import pytest

from arabic_agent.core.errors import ValidationError
from arabic_agent.documents.loader import inspect_archive, load_document, scan_for_secrets
from arabic_agent.documents.parsers import ParseLimits, supported_formats
from arabic_agent.documents.retrieval import LexicalRetriever

pytestmark = pytest.mark.asyncio


# --- helpers to build real documents ---------------------------------------


def _make_pdf(path: Path, lines: list[str]) -> None:
    """Write a minimal but valid single-page PDF with extractable text."""
    text_ops = "BT /F1 18 Tf 72 720 Td 22 TL\n"
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        text_ops += f"({escaped}) Tj T*\n"
    text_ops += "ET"
    stream = text_ops.encode("latin-1", errors="replace")

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{index} 0 obj\n".encode())
        out.write(body)
        out.write(b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
        ).encode()
    )
    path.write_bytes(out.getvalue())


def _make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("تقرير المشروع", level=1)
    document.add_paragraph("هذا مستند عربي لاختبار الاستيعاب.")
    document.add_paragraph("Section: requirements and scope.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "بند"
    table.cell(0, 1).text = "قيمة"
    table.cell(1, 0).text = "الميزانية"
    table.cell(1, 1).text = "1000"
    document.save(str(path))


def _make_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "المبيعات"
    sheet.append(["المنتج", "الكمية", "السعر"])
    sheet.append(["تفاح", 10, 5.5])
    sheet.append(["موز", 20, 3.0])
    second = workbook.create_sheet("Notes")
    second.append(["note", "value"])
    workbook.save(str(path))


def _make_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "خطة العمل"
    slide.placeholders[1].text = "الهدف الأول\nالهدف الثاني"  # noqa: RUF001
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide2.shapes.title.text = "Timeline"
    slide2.placeholders[1].text = "Q1, Q2"
    presentation.save(str(path))


# --- formats ----------------------------------------------------------------


async def test_txt_and_markdown(tmp_path: Path) -> None:
    txt = tmp_path / "notes.txt"
    txt.write_text("سطر أول\nسطر ثان\n", encoding="utf-8")
    doc = load_document(txt, display_path="notes.txt")
    assert doc.has_text
    assert doc.format == "text"
    assert "سطر أول" in doc.sections[0].text

    md = tmp_path / "readme.md"
    md.write_text("# العنوان\n\nنص الفقرة\n\n## قسم\nمحتوى\n", encoding="utf-8")
    doc = load_document(md, display_path="readme.md")
    headings = [s for s in doc.sections if s.kind == "heading"]
    assert len(headings) >= 2
    assert headings[0].level == 1
    assert headings[0].provenance.locator.startswith("lines")


async def test_json(tmp_path: Path) -> None:
    path = tmp_path / "data.json"
    path.write_text('{"name": "تقرير", "items": [1, 2, 3]}', encoding="utf-8")
    doc = load_document(path, display_path="data.json")
    assert doc.format == "json"
    assert doc.has_text
    assert "تقرير" in "".join(s.text for s in doc.sections)


async def test_json_invalid_reports_honestly(tmp_path: Path) -> None:
    path = tmp_path / "bad.json"
    path.write_text("{not valid", encoding="utf-8")
    with pytest.raises(ValidationError):
        load_document(path, display_path="bad.json")


async def test_csv_delimiter_and_rows(tmp_path: Path) -> None:
    path = tmp_path / "sales.csv"
    path.write_text("product,qty\napple,10\nbanana,20\n", encoding="utf-8")
    doc = load_document(path, display_path="sales.csv")
    assert doc.format == "csv"
    assert doc.metadata.row_count == 3
    assert doc.metadata.column_count == 2
    assert doc.metadata.extra["delimiter"] == ","
    assert any("apple" in s.text for s in doc.sections)


async def test_csv_semicolon_delimiter(tmp_path: Path) -> None:
    path = tmp_path / "eu.csv"
    path.write_text("a;b;c\n1;2;3\n", encoding="utf-8")
    doc = load_document(path, display_path="eu.csv")
    assert doc.metadata.extra["delimiter"] == ";"


@pytest.mark.skipif(
    "html" not in supported_formats(), reason="beautifulsoup4 (documents extra) not installed"
)
async def test_html(tmp_path: Path) -> None:
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>صفحة</title></head><body><h1>عنوان</h1><p>فقرة نصية</p></body></html>",
        encoding="utf-8",
    )
    doc = load_document(path, display_path="page.html")
    assert doc.format == "html"
    assert doc.metadata.title == "صفحة"
    assert any(s.kind == "heading" and "عنوان" in s.text for s in doc.sections)


@pytest.mark.skipif("pdf" not in supported_formats(), reason="pypdf not installed")
async def test_pdf_text_and_page_provenance(tmp_path: Path) -> None:
    path = tmp_path / "doc.pdf"
    _make_pdf(path, ["Hello PDF report", "Second line with data"])
    doc = load_document(path, display_path="doc.pdf")
    assert doc.format == "pdf"
    assert doc.has_text
    assert doc.metadata.page_count == 1
    joined = "".join(s.text for s in doc.sections)
    assert "Hello PDF report" in joined
    assert doc.sections[0].provenance.page == 1
    assert doc.sections[0].provenance.locator == "page 1"


@pytest.mark.skipif("pdf" not in supported_formats(), reason="pypdf not installed")
async def test_pdf_without_text_is_reported_not_faked(tmp_path: Path) -> None:
    # A PDF whose page has no text content stream.
    path = tmp_path / "blank.pdf"
    _make_pdf(path, [])
    doc = load_document(path, display_path="blank.pdf")
    assert doc.has_text is False
    assert any("no extractable text" in w for w in doc.warnings)


@pytest.mark.skipif("docx" not in supported_formats(), reason="python-docx not installed")
async def test_docx_structure_and_arabic(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    _make_docx(path)
    doc = load_document(path, display_path="report.docx")
    assert doc.format == "docx"
    headings = [s for s in doc.sections if s.kind == "heading"]
    tables = [s for s in doc.sections if s.kind == "table"]
    assert headings and headings[0].text == "تقرير المشروع"
    assert tables and "الميزانية" in tables[0].text
    assert doc.metadata.paragraph_count and doc.metadata.paragraph_count >= 2


@pytest.mark.skipif("xlsx" not in supported_formats(), reason="openpyxl not installed")
async def test_xlsx_sheets_and_ranges(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    _make_xlsx(path)
    doc = load_document(path, display_path="book.xlsx")
    assert doc.format == "xlsx"
    assert doc.metadata.sheet_count == 2
    assert "المبيعات" in doc.metadata.extra["sheets"]
    sheets = [s for s in doc.sections if s.kind == "sheet"]
    assert any(s.provenance.sheet == "المبيعات" for s in sheets)
    assert any("تفاح" in s.text for s in sheets)
    assert any(s.provenance.range for s in sheets)


@pytest.mark.skipif("pptx" not in supported_formats(), reason="python-pptx not installed")
async def test_pptx_slides_and_provenance(tmp_path: Path) -> None:
    # python-pptx imports Pillow itself, so a build that has pptx but not
    # Pillow cannot run this test either; skip honestly rather than error.
    pytest.importorskip("PIL")
    path = tmp_path / "deck.pptx"
    _make_pptx(path)
    doc = load_document(path, display_path="deck.pptx")
    assert doc.format == "pptx"
    assert doc.metadata.slide_count == 2
    slides = [s for s in doc.sections if s.kind == "slide"]
    assert len(slides) == 2
    assert slides[0].provenance.slide == 1
    assert "خطة العمل" in slides[0].text


async def test_unknown_format_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "weird.xyz"
    path.write_text("data", encoding="utf-8")
    with pytest.raises(ValidationError):
        load_document(path, display_path="weird.xyz")


# --- retrieval & provenance -------------------------------------------------


async def test_retrieval_ranks_relevant_sections(tmp_path: Path) -> None:
    path = tmp_path / "manual.txt"
    body = "\n".join(
        [f"line {i} about nothing in particular" for i in range(200)]
        + ["the warranty period is 24 months for the device"]
        + [f"trailing line {i}" for i in range(50)]
    )
    path.write_text(body, encoding="utf-8")
    doc = load_document(path, display_path="manual.txt")
    result = LexicalRetriever().retrieve(doc, "warranty period months", limit=3)
    assert result.matches
    top = result.matches[0].section.text
    assert "warranty" in top
    assert result.matches[0].section.provenance.path == "manual.txt"
    assert result.total_sections >= 1


async def test_retrieval_reports_omissions(tmp_path: Path) -> None:
    path = tmp_path / "many.txt"
    path.write_text("\n".join(f"keyword line {i}" for i in range(400)), encoding="utf-8")
    doc = load_document(path, display_path="many.txt")
    result = LexicalRetriever().retrieve(doc, "keyword", limit=2, max_tokens=10)
    assert len(result.matches) <= 2
    assert result.truncated is True


# --- archives ---------------------------------------------------------------


async def test_inspect_zip_flags_traversal(tmp_path: Path) -> None:
    path = tmp_path / "evil.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ok.txt", "fine")
        archive.writestr("../escape.txt", "bad")
    inspection = inspect_archive(path, display_path="evil.zip")
    assert "../escape.txt" in inspection.unsafe_entries
    assert inspection.format == "zip"


async def test_inspect_tar_lists_entries(tmp_path: Path) -> None:
    path = tmp_path / "bundle.tar"
    payload = tmp_path / "a.txt"
    payload.write_text("hello", encoding="utf-8")
    with tarfile.open(path, "w") as archive:
        archive.add(payload, arcname="a.txt")
    inspection = inspect_archive(path, display_path="bundle.tar")
    assert inspection.format == "tar"
    assert any(e.name == "a.txt" for e in inspection.entries)


# --- secret detection -------------------------------------------------------


async def test_secret_scan_reports_kind_not_value(tmp_path: Path) -> None:
    path = tmp_path / "config.txt"
    path.write_text("api_key = sk-abcdefghijklmnopqrstuvwx\n", encoding="utf-8")
    doc = load_document(path, display_path="config.txt")
    findings = scan_for_secrets(doc)
    assert findings
    assert any(f.kind == "openai_key" for f in findings)
    # The finding must not contain the secret value itself.
    assert all("sk-abcdefghijklmnop" not in f.kind for f in findings)


async def test_parse_limits_bound_sections(tmp_path: Path) -> None:
    path = tmp_path / "huge.txt"
    path.write_text("\n".join(f"row {i}" for i in range(5000)), encoding="utf-8")
    doc = load_document(path, display_path="huge.txt", limits=ParseLimits(max_sections=3))
    assert len(doc.sections) <= 3
    assert doc.truncated is True
